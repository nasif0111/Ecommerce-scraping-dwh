from pathlib import Path

import pandas as pd
import pyodbc
from openpyxl.utils import get_column_letter


# =========================
# Configuration
# =========================

OUTPUT_FILE = Path(
    r"G:\data clean\scrabing\jumia_products_from_sql.xlsx"
)

CONNECTION_STRING = (
    "DRIVER={ODBC Driver 18 for SQL Server};"
    r"SERVER=.\SQLEXPRESS;"
    "DATABASE=E_commerce_scrabingproj;"
    "Trusted_Connection=yes;"
    "TrustServerCertificate=yes;"
)


QUERY = """
SELECT
    ProductID,
    SKU,
    ProductName,
    ImageURL,
    ProductURL,
    CurrentPrice,
    OldPrice,
    DiscountPercent,
    Rating,
    ReviewsCount,
    st_date,
    end_date,
    iscurrent
FROM dbo.JumiaProducts
WHERE iscurrent = 1
ORDER BY ProductID;
"""


def main():
    # 1) Read data from SQL Server
    with pyodbc.connect(CONNECTION_STRING) as connection:
        df = pd.read_sql_query(QUERY, connection)

    # 2) Stop if main table has no current data
    if df.empty:
        print("No current data found in dbo.JumiaProducts.")
        return

    # Ensure output folder exists
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

    # 3) Create Excel file
    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            sheet_name="Jumia Products",
            index=False
        )

        worksheet = writer.sheets["Jumia Products"]

        # 4) Set width for every column
        for column_cells in worksheet.columns:
            column_letter = get_column_letter(
                column_cells[0].column
            )

            max_length = max(
                len(str(cell.value))
                if cell.value is not None
                else 0
                for cell in column_cells
            )

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 50)

        # 5) Date columns format and width
        for column_name in ["st_date", "end_date"]:
            if column_name in df.columns:
                column_index = (
                    df.columns.get_loc(column_name) + 1
                )

                column_letter = get_column_letter(
                    column_index
                )

                worksheet.column_dimensions[
                    column_letter
                ].width = 22

                for cell in worksheet[column_letter][1:]:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"

    print("Export completed successfully.")
    print(f"Rows exported: {len(df)}")
    print(f"Excel file created: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()